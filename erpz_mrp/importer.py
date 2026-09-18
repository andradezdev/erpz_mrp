def _bulk_update_items_sql(items_data):
    if not items_data:
        return
    query = """
        UPDATE `tabItem` SET
            item_name = %s,
            item_group = %s,
            stock_uom = %s,
            custom_mrp_reorder_point = %s,
            custom_mrp_safety_stock = %s,
            custom_mrp_min_stock = %s,
            custom_mrp_total_safety_threshold = %s,
            min_order_qty = %s,
            modified = %s,
            modified_by = %s
        WHERE name = %s
    """
    for row in items_data:
        frappe.db.sql(query, row)

# Copyright (c) 2026, ERPZ and contributors
# For license information, please see license.txt

import io
import csv
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import frappe
from frappe import _
from frappe.utils import flt, cint, getdate, nowdate, now_datetime

# -------------------------------------------------------------------------
# STYLING UTILITIES FOR ENTERPRISE EXCEL TEMPLATES
# -------------------------------------------------------------------------

def style_excel_sheet(ws, title, headers, rows, header_color="1F4E79", title_color="102A43"):
    """Applies professional enterprise styling to an openpyxl worksheet."""
    header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E0'),
        right=Side(style='thin', color='CBD5E0'),
        top=Side(style='thin', color='CBD5E0'),
        bottom=Side(style='thin', color='CBD5E0')
    )

    # Title Banner
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = title
    title_cell.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color=title_color, end_color=title_color, fill_type="solid")
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 28

    # Spacer
    ws.append([])
    ws.row_dimensions[2].height = 6

    # Headers
    ws.append(headers)
    ws.row_dimensions[3].height = 24

    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=3, column=col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border

    # Data Rows
    current_row = 4
    for r in rows:
        ws.append(r)
        fill = zebra_fill if (current_row % 2 == 0) else white_fill
        ws.row_dimensions[current_row].height = 20

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.fill = fill
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)

            val = cell.value
            if isinstance(val, (int, float)):
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        current_row += 1

    # Auto Column Widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row > 2 and cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

# -------------------------------------------------------------------------
# TEMPLATE DEFINITIONS (ALL 7 CADASTROS)
# -------------------------------------------------------------------------

TEMPLATES_DEF = {
    "items": {
        "title": "ERPZ MRP — Modelo de Importação de Produtos / Itens",
        "filename": "Modelo_Importacao_Produtos_Itens",
        "header_color": "1F4E79",
        "headers": [
            "Código do Item*", "Descrição do Produto*", "Unidade de Medida (UOM)*",
            "Grupo de Itens", "Ponto de Pedido (MRP)", "Estoque de Segurança (MRP)",
            "Estoque Mínimo (MRP)", "Lote Mínimo", "Depósito Padrão"
        ],
        "rows": [
            ["100100001", "BROCHADEIRA VERTICAL INTERNA BVI", "PC", "8626", 10.0, 5.0, 15.0, 1.0, "01"],
            ["200110016", "CHAPA DE ACO CARBONO 1020 3MM", "KG", "Materia Prima", 100.0, 50.0, 150.0, 25.0, "01"],
            ["802001929", "TUBO COM GRAXEIRA ROSQUEADA", "PC", "Produtos Acabados", 20.0, 10.0, 30.0, 5.0, "01"]
        ]
    },
    "stock": {
        "title": "ERPZ MRP — Modelo de Importação de Estoque Inicial dos Itens",
        "filename": "Modelo_Importacao_Estoque_Itens",
        "header_color": "2B6CB0",
        "headers": [
            "Código do Item*", "Depósito / Armazém*", "Quantidade em Estoque*",
            "Custo Unitário / Valor de Avaliação", "Data do Saldo (DD/MM/AAAA)", "Empresa"
        ],
        "rows": [
            ["200110016", "01", 702.35, 7.74, "30/04/2026", "Andradez Dev"],
            ["802001929", "01", 2426.0, 15.50, "30/04/2026", "Andradez Dev"],
            ["100100001", "01", 12.0, 1250.00, "30/04/2026", "Andradez Dev"]
        ]
    },
    "workstations": {
        "title": "ERPZ MRP / APS — Modelo de Importação de Recursos e Postos de Trabalho",
        "filename": "Modelo_Importacao_Recursos_Postos_Trabalho",
        "header_color": "2C5282",
        "headers": [
            "Código do Recurso / Posto*", "Nome do Recurso / Descrição*", "Empresa",
            "Eficiência Efetiva (%)", "Horas Diárias Disponíveis", "Trabalho em Finais de Semana (0 ou 1)"
        ],
        "rows": [
            ["PHC03", "PRENSA HIDRAULICA C03", "Andradez Dev", 100.0, 8.0, 0],
            ["LHR01", "LAMINADORA DE ROSCA 01", "Andradez Dev", 95.0, 8.0, 0],
            ["CUC07", "CENTRO DE USINAGEM CNC 07", "Andradez Dev", 100.0, 16.0, 1]
        ]
    },
    "alternative_resources": {
        "title": "ERPZ APS — Modelo de Importação de Recursos Alternativos",
        "filename": "Modelo_Importacao_Recursos_Alternativos",
        "header_color": "2A4365",
        "headers": [
            "Recurso Principal*", "Recurso Alternativo*", "Ordem / Prioridade", "Fator de Eficiência (%)"
        ],
        "rows": [
            ["MOE01", "MOE02", 1, 100.0],
            ["MOE01", "MOE03", 2, 90.0],
            ["PHH03", "PHH01", 1, 100.0]
        ]
    },
    "bom": {
        "title": "ERPZ MRP — Modelo de Importação de Estrutura de Produtos (BOM)",
        "filename": "Modelo_Importacao_Estrutura_BOM",
        "header_color": "1A365D",
        "headers": [
            "Código do Produto Pai (Dono do Registro)*", "Descrição do Produto Pai",
            "Código do Componente (Filho)*", "Descrição do Componente",
            "Quantidade do Componente*", "Índice de Perda (%)", "Sequência"
        ],
        "rows": [
            ["802001929", "TUBO COM GRAXEIRA ROSQUEADA", "220100015", "GRAXEIRA RETA 1/4 UNF", 1.0, 0.0, 1],
            ["802001929", "TUBO COM GRAXEIRA ROSQUEADA", "805001929", "TUBO FRESADO E ROSCADO", 1.0, 0.0, 2],
            ["802029092", "ESFERA ZINCADA", "810029092", "ESFERA TEMPERADA", 1.0, 0.0, 1]
        ]
    },
    "customers": {
        "title": "ERPZ — Modelo de Importação de Clientes",
        "filename": "Modelo_Importacao_Clientes",
        "header_color": "0D3B66",
        "headers": [
            "Código do Cliente*", "Nome do Cliente / Razão Social*", "Nome Fantasia",
            "Tipo (Física ou Jurídica)", "CNPJ / CPF", "Inscrição Estadual",
            "E-Mail", "Telefone", "Endereço", "Bairro", "Município / Cidade",
            "Estado (UF)", "CEP"
        ],
        "rows": [
            ["011829", "POLO COMERCIO DE PECAS LTDA", "POLO AGRICOLA", "Jurídica", "13.219.061/0001-96", "10.493.386-0", "poloagricola@example.com", "61-36422195", "RUA 04 44", "JARDIM CALIFORNIA", "FORMOSA", "GO", "73801-971"],
            ["000001", "CLIENTE PADRAO P/ ORCAMENTO", "ORCAMENTO", "Jurídica", "", "", "orcamento@example.com", "", "AVENIDA CENTRAL 100", "CENTRO", "SAO PAULO", "SP", "01001-000"]
        ]
    },
    "suppliers": {
        "title": "ERPZ — Modelo de Importação de Fornecedores",
        "filename": "Modelo_Importacao_Fornecedores",
        "header_color": "1D3557",
        "headers": [
            "Código do Fornecedor*", "Razão Social / Nome do Fornecedor*", "Nome Fantasia",
            "Tipo (Física ou Jurídica)", "CNPJ / CPF", "Inscrição Estadual",
            "E-Mail", "Telefone", "Endereço", "Bairro", "Município / Cidade",
            "Estado (UF)", "CEP"
        ],
        "rows": [
            ["017604", "GRASSI E GRASSI DISTR. DE LUBRIFICANTES", "GRASSI", "Jurídica", "04.870.154/0001-95", "647402379119", "contato@grassi.com.br", "17-32280-101", "RUA ULISSES JAMIL CURY", "D.ULISSES GUIMARAES", "SAO JOSE DO RIO PRETO", "SP", "15092-601"],
            ["017605", "SUL DISTRIBUIDORA DE COMP.ELETRICOS LTDA", "SUL DISTRIBUIDORA", "Jurídica", "54.052.006/0001-31", "111211072118", "vendas@suldistribuidora.com.br", "11-56417-288", "RUA AGUAI", "VILA CRUZEIRO", "SAO PAULO", "SP", "04728-030"]
        ]
    }
}

def generate_template_workbook(import_type):
    """Creates a styled openpyxl Workbook for the requested import type."""
    meta = TEMPLATES_DEF.get(import_type)
    if not meta:
        frappe.throw(_("Tipo de importação desconhecido: {0}").format(import_type))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = import_type[:31]

    style_excel_sheet(
        ws,
        title=meta["title"],
        headers=meta["headers"],
        rows=meta["rows"],
        header_color=meta.get("header_color", "1F4E79")
    )
    return wb, meta["filename"]

# -------------------------------------------------------------------------
# GENERIC ROBUST FILE READER (EXCEL + CSV COM LATIN-1 / UTF-8 / AUTO-HEADER)
# -------------------------------------------------------------------------

def clean_val(val):
    if val is None:
        return ""
    if isinstance(val, (int, float)):
        return val
    s = str(val).strip()
    if (s.startswith('"') and s.endswith('"')) or (s.startswith("'") and s.endswith("'")):
        s = s[1:-1].strip()
    return s

def sanitize_email(email_str):
    if not email_str:
        return ""
    s = str(email_str).strip()
    # Remove trailing dot before @ if present (e.g. name.@gmail.com -> name@gmail.com)
    s = re.sub(r'\.+@', '@', s)
    if re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", s):
        return s
    return ""

def parse_num(val, default=0.0):
    if val is None or val == "":
        return default
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if (s.startswith('"') and s.endswith('"')) or (s.startswith("'") and s.endswith("'")):
        s = s[1:-1].strip()
    if not s:
        return default
    s = s.replace(" ", "")
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except Exception:
        return default

def parse_date_str(val):
    if not val:
        return None
    if hasattr(val, "strftime"):
        return getdate(val)
    s = str(val).strip()
    if (s.startswith('"') and s.endswith('"')) or (s.startswith("'") and s.endswith("'")):
        s = s[1:-1].strip()
    if not s or "/" not in s or s in ("  /  /    ", ".  ."):
        return None
    try:
        parts = s.split("/")
        if len(parts) == 3 and len(parts[2]) == 4:
            return getdate(f"{parts[2]}-{parts[1]}-{parts[0]}")
        return getdate(s)
    except Exception:
        return None

def normalize_key(s):
    if not s:
        return ""
    s = str(s).lower().strip()
    s = re.sub(r'[\*\(\)\[\]\:\.\,\/\-\\\'\"]', ' ', s)
    s = re.sub(r'\s+', ' ', s)
    import unicodedata
    s = ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')
    return s.strip()

FIELD_ALIASES = {
    "items": {
        "item_code": ["codigo", "cod", "cod item", "cod interno", "codigo do item", "item code", "produto"],
        "item_name": ["descricao", "descricao do produto", "desc prod", "item name", "nome"],
        "stock_uom": ["unidade", "unidade de medida", "stock uom", "uom", "um", "unid"],
        "item_group": ["grupo", "grupo de itens", "item group", "familia"],
        "reorder_point": ["ponto pedido", "ponto de pedido", "custom mrp reorder point"],
        "safety_stock": ["seguranca", "estoque de seguranca", "custom mrp safety stock"],
        "min_stock": ["estoque minimo", "estoque a manter", "min stock", "custom mrp min stock"],
        "min_order_qty": ["lote minimo", "lote econom", "min order qty"],
        "default_warehouse": ["armazem pad", "deposito padrao", "default warehouse"]
    },
    "stock": {
        "item_code": ["produto", "codigo do item", "item code", "codigo", "item"],
        "warehouse": ["armazem", "deposito", "deposito armazem", "warehouse"],
        "qty": ["qtd inic mes", "quantidade em estoque", "quantidade", "qty", "qtd"],
        "valuation_rate": ["c unit 1a m", "custo stand", "custo unitario", "custo", "valuation rate", "preco"],
        "posting_date": ["data saldo", "data do saldo", "data", "posting date"],
        "company": ["filial", "empresa", "company"]
    },
    "workstations": {
        "workstation": ["codigo do recurso", "codigo do posto", "workstation", "recurso", "posto", "codigo"],
        "resource_name": ["nome recurso", "nome do recurso", "resource name", "descricao", "nome"],
        "company": ["filial", "empresa", "company"],
        "efficiency_factor": ["m o efic", "eficiencia efetiva", "efficiency factor", "eficiencia"],
        "capacity_hours_per_day": ["horas diarias disponiveis", "horas diarias", "capacity hours", "horas"],
        "allow_weekend_work": ["trabalho em finais de semana", "final de semana", "allow weekend work"]
    },
    "alternative_resources": {
        "primary_workstation": ["recur princ", "recurso principal", "primary workstation", "principal"],
        "alternative_workstation": ["rec alt sec", "recurso alternativo", "alternative workstation", "alternativo"],
        "priority": ["ordem", "prioridade", "priority"],
        "efficiency_factor": ["fator de eficiencia", "efficiency factor", "eficiencia"]
    },
    "bom": {
        "parent_item": ["codigo do produto pai", "dono do registro", "parent item", "codigo", "produto", "pai", "dono"],
        "parent_name": ["descricao do produto pai", "desc prod", "descricao pai"],
        "child_item": ["codigo do componente", "componente", "cod componente", "child item", "filho"],
        "child_name": ["descricao do componente", "desc comp", "descricao"],
        "qty": ["quantidade do componente", "quantidade", "qty", "qtd"],
        "scrap_percentage": ["indice perda", "scrap percentage", "perda"],
        "sequence": ["sequencia", "sequence"]
    },
    "customers": {
        "customer_code": ["codigo do cliente", "customer code", "codigo"],
        "customer_name": ["nome do cliente", "razao social", "customer name", "nome"],
        "customer_short_name": ["n fantasia", "nome fantasia", "fantasia"],
        "customer_type": ["fisica jurid", "tipo pessoa", "customer type"],
        "tax_id": ["cnpj cpf", "cnpj", "cpf", "tax id"],
        "ie": ["ins estad", "inscricao estadual"],
        "email_id": ["e mail", "email", "email id"],
        "phone": ["telefone", "fone", "mobile no", "ddd"],
        "address": ["endereco", "logradouro"],
        "neighborhood": ["bairro"],
        "city": ["municipio", "cidade"],
        "state": ["estado", "uf"],
        "pincode": ["cep"]
    },
    "suppliers": {
        "supplier_code": ["codigo do fornecedor", "supplier code", "codigo"],
        "supplier_name": ["nome do fornecedor", "razao social", "supplier name", "nome"],
        "supplier_short_name": ["n fantasia", "nome fantasia", "fantasia"],
        "supplier_type": ["tipo pessoa", "supplier type"],
        "tax_id": ["cnpj cpf", "cnpj", "cpf", "tax id"],
        "ie": ["ins estad", "inscricao estadual"],
        "email_id": ["e mail", "email", "email id"],
        "phone": ["telefone", "fone", "mobile no", "ddd"],
        "address": ["endereco", "logradouro"],
        "neighborhood": ["bairro"],
        "city": ["municipio", "cidade"],
        "state": ["estado", "uf"],
        "pincode": ["cep"]
    }
}

def parse_file_content(file_content, filename, import_type):
    """
    Parses bytes from Excel or CSV and maps columns to canonical dictionary fields.
    Returns: list of dicts: [{"row_num": 4, "data": {...}}, ...]
    """
    aliases = FIELD_ALIASES.get(import_type, {})
    is_excel = filename.lower().endswith(('.xlsx', '.xlsm', '.xltx')) or (len(file_content) >= 4 and file_content[:4] == b'PK\x03\x04')

    raw_matrix = []
    if is_excel:
        wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True, read_only=True)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            raw_matrix.append([clean_val(c) for c in row])
    else:
        decoded_text = None
        for enc in ['utf-8-sig', 'utf-8', 'latin-1', 'cp1252', 'iso-8859-1']:
            try:
                decoded_text = file_content.decode(enc)
                break
            except Exception:
                continue

        if decoded_text is None:
            decoded_text = file_content.decode('latin-1', errors='replace')

        first_few_lines = decoded_text[:2048]
        delimiter = ';' if first_few_lines.count(';') >= first_few_lines.count(',') else ','

        reader = csv.reader(decoded_text.splitlines(), delimiter=delimiter)
        for row in reader:
            raw_matrix.append([clean_val(c) for c in row])

    if not raw_matrix:
        frappe.throw(_("O arquivo enviado está vazio."))

    header_row_idx = -1
    col_mapping = {}

    for r_idx in range(min(12, len(raw_matrix))):
        row = raw_matrix[r_idx]
        norm_row = [normalize_key(c) for c in row]

        # Pass 1: Exact matches
        current_map = {}
        matched_keys = set()
        for canon_key, alias_list in aliases.items():
            for a in alias_list:
                norm_a = normalize_key(a)
                if norm_a in norm_row:
                    c_idx = norm_row.index(norm_a)
                    current_map[canon_key] = c_idx
                    matched_keys.add(canon_key)
                    break

        # Pass 2: Partial matches for remaining keys (avoiding short collisions)
        for canon_key, alias_list in aliases.items():
            if canon_key in matched_keys:
                continue
            for a in alias_list:
                norm_a = normalize_key(a)
                if len(norm_a) < 4:
                    continue
                found = False
                for c_idx, cell_norm in enumerate(norm_row):
                    if c_idx in current_map.values():
                        continue
                    if norm_a == cell_norm or (len(norm_a) >= 5 and norm_a in cell_norm):
                        current_map[canon_key] = c_idx
                        matched_keys.add(canon_key)
                        found = True
                        break
                if found:
                    break

        min_matches = 1 if import_type == "alternative_resources" else 2
        if len(matched_keys) >= min_matches:
            header_row_idx = r_idx
            col_mapping = current_map
            break

    if header_row_idx == -1:
        frappe.throw(_("Não foi possível identificar os cabeçalhos das colunas correspondentes ao cadastro '{0}'. Verifique o arquivo ou baixe o modelo padrão.").format(import_type))

    parsed_rows = []
    for r_idx in range(header_row_idx + 1, len(raw_matrix)):
        row = raw_matrix[r_idx]
        if not any(row):
            continue
        row_dict = {}
        for canon_key, col_idx in col_mapping.items():
            if col_idx < len(row):
                row_dict[canon_key] = row[col_idx]
            else:
                row_dict[canon_key] = ""
        parsed_rows.append({
            "row_num": r_idx + 1,
            "data": row_dict
        })

    return parsed_rows

# -------------------------------------------------------------------------
# IMPORT PROCESSORS FOR EACH CADASTRO
# -------------------------------------------------------------------------

def ensure_uom(uom_name):
    if not uom_name:
        return "Unit"
    uom_name = uom_name.strip().upper()
    if not frappe.db.exists("UOM", uom_name):
        u = frappe.new_doc("UOM")
        u.uom_name = uom_name
        u.insert(ignore_permissions=True)
    return uom_name

def ensure_item_group(group_name):
    if not group_name:
        group_name = "Todos os Grupos de Itens"
    group_name = group_name.strip()
    if not frappe.db.exists("Item Group", group_name):
        parent_group = "Todos os Grupos de Itens" if frappe.db.exists("Item Group", "Todos os Grupos de Itens") else "All Item Groups"
        ig = frappe.new_doc("Item Group")
        ig.item_group_name = group_name
        ig.parent_item_group = parent_group
        ig.is_group = 0
        ig.insert(ignore_permissions=True)
    return group_name

def get_default_customer_group():
    cg = frappe.db.get_value("Customer Group", {"is_group": 0}, "name")
    if cg:
        return cg
    parent = frappe.db.get_value("Customer Group", {"is_group": 1}, "name") or "All Customer Groups"
    doc = frappe.new_doc("Customer Group")
    doc.customer_group_name = "Clientes Gerais"
    doc.parent_customer_group = parent
    doc.is_group = 0
    doc.insert(ignore_permissions=True)
    return doc.name

def get_default_supplier_group():
    sg = frappe.db.get_value("Supplier Group", {"is_group": 0}, "name")
    if sg:
        return sg
    parent = frappe.db.get_value("Supplier Group", {"is_group": 1}, "name") or "All Supplier Groups"
    doc = frappe.new_doc("Supplier Group")
    doc.supplier_group_name = "Fornecedores Gerais"
    doc.parent_supplier_group = parent
    doc.is_group = 0
    doc.insert(ignore_permissions=True)
    return doc.name

def resolve_company_warehouse(company, warehouse_code):
    """Resolves or creates a valid non-group Warehouse for the given code & company."""
    if warehouse_code and frappe.db.exists("Warehouse", warehouse_code):
        return warehouse_code

    if warehouse_code:
        wh = frappe.db.get_value("Warehouse", {"company": company, "is_group": 0, "warehouse_name": ["like", f"%{warehouse_code}%"]}, "name")
        if wh:
            return wh

    wh = frappe.db.get_value("Warehouse", {"company": company, "is_group": 0}, "name")
    if wh:
        return wh

    parent_wh = frappe.db.get_value("Warehouse", {"company": company, "is_group": 1}, "name")
    wh_doc = frappe.new_doc("Warehouse")
    wh_doc.warehouse_name = f"Armazém Principal {company}"
    wh_doc.company = company
    wh_doc.is_group = 0
    if parent_wh:
        wh_doc.parent_warehouse = parent_wh
    wh_doc.insert(ignore_permissions=True)
    return wh_doc.name

def resolve_company_opening_account(company):
    """Finds or creates a Balance Sheet Temporary Opening Account for Stock Reconciliation."""
    acc = frappe.db.get_value("Account", {"company": company, "account_type": "Temporary", "is_group": 0}, "name")
    if acc:
        return acc
    acc = frappe.db.get_value("Account", {"company": company, "name": ["like", "%Abertura%"], "is_group": 0, "report_type": "Balance Sheet"}, "name")
    if acc:
        return acc

    parent_liab = frappe.db.get_value("Account", {"company": company, "root_type": "Liability", "is_group": 1}, "name")
    if not parent_liab:
        parent_liab = frappe.db.get_value("Account", {"company": company, "is_group": 1}, "name")
    if parent_liab:
        doc = frappe.new_doc("Account")
        doc.account_name = "Conta Transitória de Abertura"
        doc.company = company
        doc.parent_account = parent_liab
        doc.account_type = "Temporary"
        doc.insert(ignore_permissions=True)
        return doc.name
    return None

# 1. IMPORT PRODUTOS / ITENS (SB1)
def process_import_items(rows, company=None):
    created = 0
    updated = 0
    errors = []
    total = len(rows)

    # 1. Pré-cache em memória de alta performance
    existing_items = set(frappe.db.sql_list("SELECT name FROM tabItem"))
    existing_uoms = set(frappe.db.sql_list("SELECT name FROM tabUOM"))
    existing_groups = set(frappe.db.sql_list("SELECT name FROM `tabItem Group`"))

    default_group = "Todos os Grupos de Itens"
    if default_group not in existing_groups:
        default_group = ensure_item_group(default_group)
        existing_groups.add(default_group)

    item_fields = [
        "name", "item_code", "item_name", "item_group", "stock_uom", "is_stock_item",
        "custom_mrp_reorder_point", "custom_mrp_safety_stock", "custom_mrp_min_stock",
        "custom_mrp_total_safety_threshold", "min_order_qty",
        "creation", "modified", "owner", "modified_by"
    ]

    items_to_insert = []
    items_to_update = []
    now_dt = now_datetime()
    user = frappe.session.user or "Administrator"

    for idx, item_row in enumerate(rows, 1):
        d = item_row["data"]
        code = clean_val(d.get("item_code"))
        if not code:
            continue

        name = clean_val(d.get("item_name")) or code
        uom_raw = clean_val(d.get("stock_uom")) or "PC"
        uom = uom_raw.strip().upper()
        if uom not in existing_uoms:
            ensure_uom(uom)
            existing_uoms.add(uom)

        group_raw = clean_val(d.get("item_group")) or default_group
        if group_raw not in existing_groups:
            ensure_item_group(group_raw)
            existing_groups.add(group_raw)
        group = group_raw

        reorder = parse_num(d.get("reorder_point"), 0.0)
        safety = parse_num(d.get("safety_stock"), 0.0)
        min_stock = parse_num(d.get("min_stock"), 0.0)
        min_order = parse_num(d.get("min_order_qty"), 0.0)
        total_safety = reorder + safety

        if code in existing_items:
            items_to_update.append((
                name, group, uom, reorder, safety, min_stock, total_safety, min_order, now_dt, user, code
            ))
            updated += 1
        else:
            items_to_insert.append([
                code, code, name, group, uom, 1, reorder, safety, min_stock, total_safety, min_order,
                now_dt, now_dt, user, user
            ])
            existing_items.add(code)
            created += 1

        # Insere em lotes ultra rápidos a cada 1.000 registros
        if len(items_to_insert) >= 1000:
            frappe.db.bulk_insert("Item", item_fields, items_to_insert)
            items_to_insert = []
            frappe.db.commit()

        # Atualiza em lotes a cada 500 registros
        if len(items_to_update) >= 500:
            _bulk_update_items_sql(items_to_update)
            items_to_update = []
            frappe.db.commit()

        if idx % 5000 == 0:
            pct = int((idx / total) * 100)
            frappe.publish_progress(pct, title="Processando Carga de Itens no ERPZ...", description=f"{idx:,} de {total:,} processados ({pct}%)")

    if items_to_insert:
        frappe.db.bulk_insert("Item", item_fields, items_to_insert)
        items_to_insert = []

    if items_to_update:
        _bulk_update_items_sql(items_to_update)
        items_to_update = []


    # Relink automático de sub-BOMs multinível
    frappe.db.sql("""
        UPDATE `tabBOM Item` bi
        INNER JOIN `tabBOM` child_bom ON child_bom.item = bi.item_code 
            AND child_bom.is_active = 1 
            AND child_bom.is_default = 1 
            AND child_bom.docstatus = 1
        SET bi.bom_no = child_bom.name
        WHERE (bi.bom_no IS NULL OR bi.bom_no = '')
    """)
    frappe.db.commit()
    return created, updated, errors


# 2. IMPORT ESTOQUE DOS ITENS (SB9)
def process_import_stock(rows, company=None):
    if not company:
        company = frappe.defaults.get_user_default("Company") or frappe.db.get_value("Company", {}, "name")

    created = 0
    updated = 0
    errors = []

    opening_acc = resolve_company_opening_account(company)
    items_to_reconcile = []

    for stock_row in rows:
        row_num = stock_row["row_num"]
        d = stock_row["data"]
        code = clean_val(d.get("item_code"))
        if not code:
            continue

        qty = parse_num(d.get("qty"), 0.0)
        if qty <= 0:
            continue

        rate = parse_num(d.get("valuation_rate"), 0.0)
        if rate <= 0:
            rate = 1.0

        wh_raw = clean_val(d.get("warehouse"))
        wh = resolve_company_warehouse(company, wh_raw)

        if not frappe.db.exists("Item", code):
            try:
                it = frappe.new_doc("Item")
                it.item_code = code
                it.item_name = code
                it.stock_uom = "PC"
                it.item_group = ensure_item_group("Todos os Grupos de Itens")
                it.is_stock_item = 1
                it.insert(ignore_permissions=True)
            except Exception as e:
                errors.append(f"Linha {row_num}: Não foi possível auto-criar item '{code}': {str(e)}")
                continue

        items_to_reconcile.append({
            "item_code": code,
            "warehouse": wh,
            "qty": qty,
            "valuation_rate": rate,
            "row_num": row_num
        })

    if not items_to_reconcile:
        return 0, 0, errors

    batch_size = 100
    for i in range(0, len(items_to_reconcile), batch_size):
        batch = items_to_reconcile[i:i + batch_size]
        try:
            sr = frappe.new_doc("Stock Reconciliation")
            sr.company = company
            sr.purpose = "Opening Stock"
            if opening_acc:
                sr.expense_account = opening_acc
            sr.posting_date = nowdate()
            sr.posting_time = "00:00:00"

            for item in batch:
                sr.append("items", {
                    "item_code": item["item_code"],
                    "warehouse": item["warehouse"],
                    "qty": item["qty"],
                    "valuation_rate": item["valuation_rate"]
                })

            sr.insert(ignore_permissions=True)
            try:
                sr.submit()
                created += len(batch)
            except Exception as sub_err:
                created += len(batch)
                errors.append(f"Lote {i//batch_size + 1}: Salvo como Rascunho ({sr.name}), motivo: {str(sub_err)}")

            frappe.db.commit()
        except Exception as e:
            errors.append(f"Erro ao gerar Reconciliação do lote {i//batch_size + 1}: {str(e)}")

    return created, updated, errors

# 3. IMPORT RECURSOS / POSTOS DE TRABALHO (SH1)
def process_import_workstations(rows, company=None):
    if not company:
        company = frappe.defaults.get_user_default("Company") or frappe.db.get_value("Company", {}, "name")

    created = 0
    updated = 0
    errors = []

    has_aps = frappe.db.exists("DocType", "APS Resource")

    for ws_row in rows:
        row_num = ws_row["row_num"]
        d = ws_row["data"]
        code = clean_val(d.get("workstation"))
        if not code:
            continue

        name = clean_val(d.get("resource_name")) or code
        raw_efic = parse_num(d.get("efficiency_factor"), 100.0)
        efic = (raw_efic * 100.0) if (raw_efic <= 1.0 and raw_efic > 0) else raw_efic
        hours = parse_num(d.get("capacity_hours_per_day"), 8.0)
        weekend_raw = clean_val(d.get("allow_weekend_work"))
        weekend = 1 if weekend_raw in (1, "1", "Sim", "sim", "S", "s", "true", "True") else 0

        try:
            if not frappe.db.exists("Workstation", code):
                ws = frappe.new_doc("Workstation")
                ws.workstation_name = code
                ws.description = name
                ws.production_capacity = 1
                ws.insert(ignore_permissions=True)
            else:
                frappe.db.set_value("Workstation", code, "description", name)

            if has_aps:
                if frappe.db.exists("APS Resource", code):
                    res = frappe.get_doc("APS Resource", code)
                    res.db_set({
                        "resource_name": name,
                        "company": company,
                        "efficiency_factor": efic,
                        "capacity_hours_per_day": hours,
                        "allow_weekend_work": weekend,
                        "is_active": 1
                    })
                    updated += 1
                else:
                    res = frappe.new_doc("APS Resource")
                    res.workstation = code
                    res.resource_name = name
                    res.company = company
                    res.efficiency_factor = efic
                    res.capacity_hours_per_day = hours
                    res.allow_weekend_work = weekend
                    res.is_active = 1
                    res.insert(ignore_permissions=True)
                    created += 1
            else:
                created += 1

            if (created + updated) % 50 == 0:
                frappe.db.commit()
        except Exception as e:
            errors.append(f"Linha {row_num} (Recurso {code}): {str(e)}")


    # Relink automático de sub-BOMs multinível
    frappe.db.sql("""
        UPDATE `tabBOM Item` bi
        INNER JOIN `tabBOM` child_bom ON child_bom.item = bi.item_code 
            AND child_bom.is_active = 1 
            AND child_bom.is_default = 1 
            AND child_bom.docstatus = 1
        SET bi.bom_no = child_bom.name
        WHERE (bi.bom_no IS NULL OR bi.bom_no = '')
    """)
    frappe.db.commit()
    return created, updated, errors


# 4. IMPORT RECURSOS ALTERNATIVOS (SH2)
def process_import_alternative_resources(rows, company=None):
    if not company:
        company = frappe.defaults.get_user_default("Company") or frappe.db.get_value("Company", {}, "name")

    created = 0
    updated = 0
    errors = []

    has_aps = frappe.db.exists("DocType", "APS Resource")

    for alt_row in rows:
        row_num = alt_row["row_num"]
        d = alt_row["data"]
        primary = clean_val(d.get("primary_workstation"))
        alt = clean_val(d.get("alternative_workstation"))
        if not primary or not alt:
            continue

        priority = cint(d.get("priority")) or 1
        raw_efic = parse_num(d.get("efficiency_factor"), 100.0)
        efic = (raw_efic * 100.0) if (raw_efic <= 1.0 and raw_efic > 0) else raw_efic

        try:
            for ws_code in [primary, alt]:
                if not frappe.db.exists("Workstation", ws_code):
                    w = frappe.new_doc("Workstation")
                    w.workstation_name = ws_code
                    w.production_capacity = 1
                    w.insert(ignore_permissions=True)

            if has_aps:
                if not frappe.db.exists("APS Resource", primary):
                    pres = frappe.new_doc("APS Resource")
                    pres.workstation = primary
                    pres.resource_name = primary
                    pres.company = company
                    pres.efficiency_factor = 100.0
                    pres.capacity_hours_per_day = 8.0
                    pres.is_active = 1
                    pres.insert(ignore_permissions=True)

                res_doc = frappe.get_doc("APS Resource", primary)
                existing = False
                for r in res_doc.alternatives:
                    if r.alternative_workstation == alt:
                        r.priority = priority
                        r.efficiency_factor = efic
                        existing = True
                        updated += 1
                        break
                if not existing:
                    res_doc.append("alternatives", {
                        "alternative_workstation": alt,
                        "priority": priority,
                        "efficiency_factor": efic
                    })
                    created += 1

                res_doc.save(ignore_permissions=True)
            else:
                created += 1
        except Exception as e:
            errors.append(f"Linha {row_num} (Principal {primary} -> Alt {alt}): {str(e)}")


    # Relink automático de sub-BOMs multinível
    frappe.db.sql("""
        UPDATE `tabBOM Item` bi
        INNER JOIN `tabBOM` child_bom ON child_bom.item = bi.item_code 
            AND child_bom.is_active = 1 
            AND child_bom.is_default = 1 
            AND child_bom.docstatus = 1
        SET bi.bom_no = child_bom.name
        WHERE (bi.bom_no IS NULL OR bi.bom_no = '')
    """)
    frappe.db.commit()
    return created, updated, errors


# 5. IMPORT ESTRUTURA BOM DOS ITENS (SG1 - Dono do registro e Componente lado a lado)
def process_import_bom(rows, company=None):
    if not company:
        company = frappe.defaults.get_user_default("Company") or frappe.db.get_value("Company", {}, "name")

    currency = frappe.db.get_value("Company", company, "default_currency") or "BRL"
    now_dt = now_datetime()
    user = frappe.session.user or "Administrator"

    from collections import defaultdict
    boms_dict = defaultdict(list)

    for bom_row in rows:
        d = bom_row["data"]
        p_code = clean_val(d.get("parent_item"))
        c_code = clean_val(d.get("child_item"))
        if not p_code or not c_code or p_code.lower().startswith("código"):
            continue

        qty = parse_num(d.get("qty"), 1.0)
        if qty <= 0:
            qty = 1.0
        scrap = parse_num(d.get("scrap_percentage"), 0.0)

        boms_dict[p_code].append({
            "child_code": c_code,
            "parent_name": clean_val(d.get("parent_name")) or p_code,
            "child_name": clean_val(d.get("child_name")) or c_code,
            "qty": qty,
            "scrap": scrap
        })

    # 1. Pré-cache e criação de Itens faltantes em lote
    existing_items = set(frappe.db.sql_list("SELECT name FROM tabItem"))
    missing_items = []
    it_fields = ["name", "item_code", "item_name", "item_group", "stock_uom", "is_stock_item", "custom_mrp_reorder_point", "custom_mrp_safety_stock", "custom_mrp_min_stock", "custom_mrp_total_safety_threshold", "min_order_qty", "creation", "modified", "owner", "modified_by"]

    for p_code, comps in boms_dict.items():
        if p_code not in existing_items:
            missing_items.append([p_code, p_code, comps[0]["parent_name"] or p_code, "Todos os Grupos de Itens", "PC", 1, 0, 0, 0, 0, 0, now_dt, now_dt, user, user])
            existing_items.add(p_code)
        for c in comps:
            c_code = c["child_code"]
            if c_code not in existing_items:
                missing_items.append([c_code, c_code, c["child_name"] or c_code, "Todos os Grupos de Itens", "PC", 1, 0, 0, 0, 0, 0, now_dt, now_dt, user, user])
                existing_items.add(c_code)

    if missing_items:
        for i in range(0, len(missing_items), 1000):
            frappe.db.bulk_insert("Item", it_fields, missing_items[i:i+1000])
        frappe.db.commit()

    # 2. Criação em lote de BOMs e Itens de BOM
    existing_boms = set(frappe.db.sql_list("SELECT item FROM tabBOM WHERE docstatus=1"))
    bom_fields = ["name", "item", "item_name", "quantity", "uom", "company", "currency", "is_active", "is_default", "docstatus", "creation", "modified", "owner", "modified_by"]
    bom_item_fields = ["name", "parent", "parenttype", "parentfield", "item_code", "item_name", "qty", "stock_qty", "uom", "stock_uom", "rate", "amount", "docstatus", "idx", "creation", "modified", "owner", "modified_by"]

    boms_to_insert = []
    bom_items_to_insert = []
    created = 0
    updated = 0
    errors = []

    for p_code, comps in boms_dict.items():
        if p_code in existing_boms:
            updated += 1
            continue

        bom_name = f"BOM-{p_code}-001"
        p_name = comps[0]["parent_name"] or p_code
        boms_to_insert.append([bom_name, p_code, p_name, 1.0, "PC", company, currency, 1, 1, 1, now_dt, now_dt, user, user])
        existing_boms.add(p_code)
        created += 1

        for idx, c in enumerate(comps, 1):
            item_row_name = frappe.generate_hash(length=10)
            c_code = c["child_code"]
            c_name = c["child_name"] or c_code
            q = c["qty"]
            bom_items_to_insert.append([item_row_name, bom_name, "BOM", "items", c_code, c_name, q, q, "PC", "PC", 0.0, 0.0, 1, idx, now_dt, now_dt, user, user])

    if boms_to_insert:
        for i in range(0, len(boms_to_insert), 1000):
            frappe.db.bulk_insert("BOM", bom_fields, boms_to_insert[i:i+1000])
        for i in range(0, len(bom_items_to_insert), 1000):
            frappe.db.bulk_insert("BOM Item", bom_item_fields, bom_items_to_insert[i:i+1000])
        frappe.db.commit()

    # 3. RELINK MULTINÍVEL AUTOMÁTICO
    frappe.db.sql("""
        UPDATE `tabBOM Item` bi
        INNER JOIN `tabBOM` child_bom ON child_bom.item = bi.item_code 
            AND child_bom.is_active = 1 
            AND child_bom.is_default = 1 
            AND child_bom.docstatus = 1
        SET bi.bom_no = child_bom.name
        WHERE (bi.bom_no IS NULL OR bi.bom_no = '')
    """)
    frappe.db.commit()

    return created, updated, errors


# 6. IMPORT CLIENTES (SA1)
def process_import_customers(rows, company=None):
    created = 0
    updated = 0
    errors = []

    cust_group = get_default_customer_group()
    territory = "Brasil" if frappe.db.exists("Territory", "Brasil") else "All Territories"

    for cust_row in rows:
        row_num = cust_row["row_num"]
        d = cust_row["data"]
        name = clean_val(d.get("customer_name"))
        code = clean_val(d.get("customer_code"))
        if not name:
            name = code
        if not name:
            continue

        tax_id = clean_val(d.get("tax_id"))
        raw_type = clean_val(d.get("customer_type")).lower()
        c_type = "Company" if ("jurid" in raw_type or len(tax_id) > 14) else "Individual"
        email = sanitize_email(clean_val(d.get("email_id")))
        phone = clean_val(d.get("phone"))

        try:
            existing_cust = None
            if frappe.db.exists("Customer", name):
                existing_cust = name
            elif tax_id and frappe.db.exists("Customer", {"tax_id": tax_id}):
                existing_cust = frappe.db.get_value("Customer", {"tax_id": tax_id}, "name")

            if existing_cust:
                c_doc = frappe.get_doc("Customer", existing_cust)
                if tax_id and not c_doc.tax_id:
                    c_doc.db_set("tax_id", tax_id)
                if email and not c_doc.email_id:
                    c_doc.db_set("email_id", email)
                if phone and not c_doc.mobile_no:
                    c_doc.db_set("mobile_no", phone)
                updated += 1
            else:
                c_doc = frappe.new_doc("Customer")
                c_doc.customer_name = name
                c_doc.customer_type = c_type
                c_doc.customer_group = cust_group
                c_doc.territory = territory
                if tax_id:
                    c_doc.tax_id = tax_id
                if email:
                    c_doc.email_id = email
                if phone:
                    c_doc.mobile_no = phone
                c_doc.insert(ignore_permissions=True)
                created += 1

            if (created + updated) % 100 == 0:
                frappe.db.commit()
        except Exception as e:
            errors.append(f"Linha {row_num} (Cliente {name}): {str(e)}")


    # Relink automático de sub-BOMs multinível
    frappe.db.sql("""
        UPDATE `tabBOM Item` bi
        INNER JOIN `tabBOM` child_bom ON child_bom.item = bi.item_code 
            AND child_bom.is_active = 1 
            AND child_bom.is_default = 1 
            AND child_bom.docstatus = 1
        SET bi.bom_no = child_bom.name
        WHERE (bi.bom_no IS NULL OR bi.bom_no = '')
    """)
    frappe.db.commit()
    return created, updated, errors


# 7. IMPORT FORNECEDORES (SA2)
def process_import_suppliers(rows, company=None):
    created = 0
    updated = 0
    errors = []

    supp_group = get_default_supplier_group()

    for supp_row in rows:
        row_num = supp_row["row_num"]
        d = supp_row["data"]
        name = clean_val(d.get("supplier_name"))
        code = clean_val(d.get("supplier_code"))
        if not name:
            name = code
        if not name:
            continue

        tax_id = clean_val(d.get("tax_id"))
        raw_type = clean_val(d.get("supplier_type")).lower()
        s_type = "Company" if ("jurid" in raw_type or len(tax_id) > 14) else "Individual"
        email = sanitize_email(clean_val(d.get("email_id")))
        phone = clean_val(d.get("phone"))

        try:
            existing_supp = None
            if frappe.db.exists("Supplier", name):
                existing_supp = name
            elif tax_id and frappe.db.exists("Supplier", {"tax_id": tax_id}):
                existing_supp = frappe.db.get_value("Supplier", {"tax_id": tax_id}, "name")

            if existing_supp:
                s_doc = frappe.get_doc("Supplier", existing_supp)
                if tax_id and not s_doc.tax_id:
                    s_doc.db_set("tax_id", tax_id)
                if email and not s_doc.email_id:
                    s_doc.db_set("email_id", email)
                if phone and not s_doc.mobile_no:
                    s_doc.db_set("mobile_no", phone)
                updated += 1
            else:
                s_doc = frappe.new_doc("Supplier")
                s_doc.supplier_name = name
                s_doc.supplier_type = s_type
                s_doc.supplier_group = supp_group
                if tax_id:
                    s_doc.tax_id = tax_id
                if email:
                    s_doc.email_id = email
                if phone:
                    s_doc.mobile_no = phone
                s_doc.country = "Brazil"
                s_doc.insert(ignore_permissions=True)
                created += 1

            if (created + updated) % 100 == 0:
                frappe.db.commit()
        except Exception as e:
            errors.append(f"Linha {row_num} (Fornecedor {name}): {str(e)}")


    # Relink automático de sub-BOMs multinível
    frappe.db.sql("""
        UPDATE `tabBOM Item` bi
        INNER JOIN `tabBOM` child_bom ON child_bom.item = bi.item_code 
            AND child_bom.is_active = 1 
            AND child_bom.is_default = 1 
            AND child_bom.docstatus = 1
        SET bi.bom_no = child_bom.name
        WHERE (bi.bom_no IS NULL OR bi.bom_no = '')
    """)
    frappe.db.commit()
    return created, updated, errors


# -------------------------------------------------------------------------
# MASTER DISPATCHER
# -------------------------------------------------------------------------

PROCESSORS = {
    "items": process_import_items,
    "stock": process_import_stock,
    "workstations": process_import_workstations,
    "alternative_resources": process_import_alternative_resources,
    "bom": process_import_bom,
    "customers": process_import_customers,
    "suppliers": process_import_suppliers
}

def execute_import(import_type, file_content, filename, company=None):
    """Unified handler to parse file and run the specific cadastro importer."""
    if import_type not in PROCESSORS:
        frappe.throw(_("Tipo de cadastro inválido para importação: {0}").format(import_type))

    rows = parse_file_content(file_content, filename, import_type)
    if not rows:
        frappe.throw(_("Nenhuma linha válida com dados foi encontrada no arquivo enviado."))

    handler = PROCESSORS[import_type]
    created, updated, errors = handler(rows, company)

    return {
        "success": True,
        "import_type": import_type,
        "total_rows": len(rows),
        "created_count": created,
        "updated_count": updated,
        "errors": errors
    }
