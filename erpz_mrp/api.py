# Copyright (c) 2026, ERPZ and contributors
# For license information, please see license.txt

import frappe
from frappe import _
from frappe.utils import now_datetime, getdate, flt, cint
from erpz_mrp.engine.mrp_engine import MRPEngine
from erpz_mrp.engine.execution import execute_ticket_abastecimento

@frappe.whitelist()
def run_mrp_calculation(ticket_name):
    """Triggers the 20-step MRP calculation engine for a given ticket."""
    if not frappe.has_permission("MRP Ticket", "write"):
        frappe.throw(_("Sem permissão para executar cálculo de MRP."))
        
    engine = MRPEngine(ticket_name)
    result = engine.run()
    return result

@frappe.whitelist()
def approve_ticket(ticket_name):
    """Approves an MRP ticket for subsequent document execution."""
    if not frappe.has_permission("MRP Ticket", "write"):
        frappe.throw(_("Sem permissão para aprovar Ticket MRP."))
        
    ticket = frappe.get_doc("MRP Ticket", ticket_name)
    if ticket.status not in ("Calculado", "Calculado (Simulação)", "Em Análise", "Com Inconsistências"):
        frappe.throw(_("Apenas tickets calculados podem ser aprovados."))
        
    ticket.db_set({
        "status": "Aprovado",
        "approved_on": now_datetime(),
        "approved_by": frappe.session.user
    })
    frappe.db.commit()
    return {"status": "success", "message": _("Ticket aprovado com sucesso.")}

@frappe.whitelist()
def execute_ticket(ticket_name, selected_results=None):
    """Generates Work Orders, Material Requests, and Stock Entries from MRP results."""
    if not frappe.has_permission("MRP Ticket", "write"):
        frappe.throw(_("Sem permissão para efetivar Ticket MRP."))
        
    if isinstance(selected_results, str):
        import json
        selected_results = json.loads(selected_results)
        
    res = execute_ticket_abastecimento(ticket_name, selected_results)
    return res

@frappe.whitelist()
def get_mrp_summary(ticket_name, company=None, item_code=None, supply_type=None, status=None, start=0, page_length=50):
    """Returns paginated MRP Results with filters for the Desk grid."""
    filters = {"mrp_ticket": ticket_name}
    if company:
        filters["company"] = company
    if supply_type:
        filters["supply_type"] = supply_type
    if status:
        filters["status"] = status
        
    or_filters = None
    if item_code:
        # Search across item_code, item_name, origin_name, or generated_docname
        or_filters = [
            ["item_code", "like", f"%{item_code}%"],
            ["item_name", "like", f"%{item_code}%"],
            ["origin_name", "like", f"%{item_code}%"],
            ["generated_docname", "like", f"%{item_code}%"]
        ]
        
    if or_filters:
        total_count = len(frappe.get_all("MRP Result", filters=filters, or_filters=or_filters, pluck="name"))
    else:
        total_count = frappe.db.count("MRP Result", filters=filters)
    results = frappe.get_all(
        "MRP Result",
        filters=filters,
        or_filters=or_filters,
        fields=[
            "name", "item_code", "item_name", "company", "warehouse",
            "need_date", "supply_date", "gross_demand", "initial_stock",
            "planned_inflows", "projected_balance", "safety_stock",
            "min_stock", "reorder_point", "critical_threshold", "shortage_type",
            "net_requirement", "suggested_qty", "supply_type",
            "origin_item", "origin_doctype", "origin_name", "bom_no",
            "bom_level", "lead_time_days", "from_company", "from_warehouse",
            "status", "situation", "generated_doctype", "generated_docname"
        ],
        start=cint(start),
        page_length=cint(page_length),
        order_by="need_date ASC, item_code ASC"
    )
    
    return {
        "total": total_count,
        "items": results
    }

@frappe.whitelist()
def get_timeline_data(ticket_name, item_code=None, company=None):
    """Returns day-by-day projected balance and chart series for an item or consolidated general view."""
    if item_code:
        filters = {"mrp_ticket": ticket_name, "item_code": item_code}
        if company:
            filters["company"] = company
            
        rows = frappe.get_all(
            "MRP Timeline",
            filters=filters,
            fields=[
                "item_code", "item_name", "company", "timeline_date",
                "initial_balance", "inflows", "outflows", "projected_balance",
                "safety_stock", "min_stock", "reorder_point", "critical_threshold",
                "shortage_qty", "suggested_inflow", "has_shortage", "shortage_type"
            ],
            order_by="timeline_date ASC"
        )
        
        dates = []
        projected = []
        safety = []
        min_stocks = []
        critical_thresholds = []
        shortages = []
        inflows = []
        outflows = []
        
        for r in rows:
            dates.append(str(r.timeline_date))
            projected.append(round(flt(r.projected_balance), 2))
            safety.append(round(flt(r.safety_stock), 2))
            min_stocks.append(round(flt(r.get("min_stock", 0.0)), 2))
            critical_thresholds.append(round(flt(r.get("critical_threshold", 0.0)), 2))
            shortages.append(round(flt(r.shortage_qty), 2))
            inflows.append(round(flt(r.inflows) + flt(r.suggested_inflow), 2))
            outflows.append(round(flt(r.outflows), 2))
            
        return {
            "is_general": False,
            "dates": dates,
            "projected_balance": projected,
            "safety_stock": safety,
            "min_stock": min_stocks,
            "critical_threshold": critical_thresholds,
            "shortage_qty": shortages,
            "inflows": inflows,
            "outflows": outflows,
            "raw_rows": rows
        }

    # Visão Geral Consolidada agrupada por data no banco de dados (evita sobrecarga do navegador)
    company_cond = "AND company = %(company)s" if company else ""
    params = {"ticket_name": ticket_name, "company": company}

    agg_rows = frappe.db.sql(f"""
        SELECT 
            timeline_date,
            SUM(initial_balance) as initial_balance,
            SUM(inflows + suggested_inflow) as inflows,
            SUM(outflows) as outflows,
            SUM(projected_balance) as projected_balance,
            SUM(safety_stock) as safety_stock,
            SUM(min_stock) as min_stock,
            SUM(critical_threshold) as critical_threshold,
            SUM(shortage_qty) as shortage_qty,
            MAX(has_shortage) as has_shortage
        FROM `tabMRP Timeline`
        WHERE mrp_ticket = %(ticket_name)s {company_cond}
        GROUP BY timeline_date
        ORDER BY timeline_date ASC
    """, params, as_dict=True)

    dates = []
    projected = []
    safety = []
    min_stocks = []
    critical_thresholds = []
    shortages = []
    inflows = []
    outflows = []

    for r in agg_rows:
        dates.append(str(r.timeline_date))
        projected.append(round(flt(r.projected_balance), 2))
        safety.append(round(flt(r.safety_stock), 2))
        min_stocks.append(round(flt(r.min_stock), 2))
        critical_thresholds.append(round(flt(r.critical_threshold), 2))
        shortages.append(round(flt(r.shortage_qty), 2))
        inflows.append(round(flt(r.inflows), 2))
        outflows.append(round(flt(r.outflows), 2))

    # Tabela: principais ocorrências de ruptura ou movimentação relevante (limite de 250 linhas para fluidez)
    table_rows = frappe.db.sql(f"""
        SELECT 
            item_code, item_name, company, timeline_date,
            initial_balance, inflows, outflows, projected_balance,
            safety_stock, min_stock, reorder_point, critical_threshold,
            shortage_qty, suggested_inflow, has_shortage, shortage_type
        FROM `tabMRP Timeline`
        WHERE mrp_ticket = %(ticket_name)s {company_cond}
          AND (has_shortage = 1 OR outflows > 0 OR inflows > 0 OR suggested_inflow > 0)
        ORDER BY has_shortage DESC, shortage_qty DESC, timeline_date ASC
        LIMIT 250
    """, params, as_dict=True)

    return {
        "is_general": True,
        "dates": dates,
        "projected_balance": projected,
        "safety_stock": safety,
        "min_stock": min_stocks,
        "critical_threshold": critical_thresholds,
        "shortage_qty": shortages,
        "inflows": inflows,
        "outflows": outflows,
        "raw_rows": table_rows
    }

@frappe.whitelist()
def get_traceability_tree(ticket_name, root_demand=None):
    """
    Constrói a árvore de rastreabilidade hierárquica fiel multinível:
    Nível Raiz: Pedido de Venda (ex: SAL-ORD-2026-00037)
      └── Nível 0: Produto Acabado (YAM-1627) - Produção
          ├── Nível 1: Subconjunto 805001627 - Produção
          │    └── Nível 2: Subconjunto 820001627 - Produção
          │         └── Nível 3: Subconjunto 825001627 - Produção
          │              └── Nível 4: Subconjunto 830001627 - Produção
          │                   └── Nível 5: Matéria-Prima 200220004 - Compra
          ├── Nível 1: Subconjunto 805031409 - Produção
          │    └── Nível 2: Subconjunto 805031032 - Produção
          │         └── Nível 3: Subconjunto 810031032 - Produção
          │              └── Nível 4: Subconjunto 815031024 - Produção
          │                   └── Nível 5: Subconjunto 820031024 - Produção
          │                        ├── Nível 6: Matéria-Prima 986400000 - Compra
          │                        └── Nível 6: Matéria-Prima 200300009 - Compra
          ├── Nível 1: Matéria-Prima 220100044 - Compra
          ├── Nível 1: Matéria-Prima 220100021 - Compra
          └── Nível 1: Matéria-Prima 220100005 - Compra
    """
    from collections import defaultdict
    from frappe.utils import cint

    nodes = frappe.get_all(
        "MRP Traceability",
        filters={"mrp_ticket": ticket_name, "allocated_qty": [">", 0]},
        fields=[
            "name", "demand_source_doctype", "demand_source_name",
            "root_item", "parent_item", "child_item", "bom_level",
            "required_qty", "allocated_qty", "supply_type",
            "from_company", "to_company", "need_date", "planned_start_date",
            "target_doctype", "target_docname", "status"
        ],
        order_by="bom_level ASC, need_date ASC"
    )

    if not nodes:
        nodes = frappe.get_all(
            "MRP Traceability",
            filters={"mrp_ticket": ticket_name},
            fields=[
                "name", "demand_source_doctype", "demand_source_name",
                "root_item", "parent_item", "child_item", "bom_level",
                "required_qty", "allocated_qty", "supply_type",
                "from_company", "to_company", "need_date", "planned_start_date",
                "target_doctype", "target_docname", "status"
            ],
            order_by="bom_level ASC, need_date ASC"
        )

    item_names_cache = {}
    def get_item_name(code):
        if not code:
            return ""
        if code not in item_names_cache:
            item_names_cache[code] = frappe.db.get_value("Item", code, "item_name") or code
        return item_names_cache[code]

    children_by_parent = defaultdict(list)
    root_nodes = []

    for n in nodes:
        n_dict = dict(n)
        n_dict["item_name"] = get_item_name(n.get("child_item"))
        n_dict["children"] = []

        lvl = cint(n.get("bom_level", 0))
        if lvl == 0:
            root_nodes.append(n_dict)
        else:
            parent = None
            if n.get("demand_source_name") and "Sugestão OP: " in n["demand_source_name"]:
                parent = n["demand_source_name"].replace("Sugestão OP: ", "").strip()
            elif n.get("parent_item") and n.get("parent_item") != n.get("root_item"):
                parent = n.get("parent_item")
            else:
                parent = n.get("root_item")
            children_by_parent[(parent, lvl)].append(n_dict)

    visited = set()
    def build_tree_fast(parent_node):
        if parent_node["name"] in visited:
            return
        visited.add(parent_node["name"])

        next_lvl = cint(parent_node["bom_level"]) + 1
        p_code = parent_node["child_item"]

        candidates = children_by_parent.get((p_code, next_lvl), [])
        for cand in candidates:
            cand_copy = dict(cand)
            cand_copy["children"] = []
            build_tree_fast(cand_copy)
            parent_node["children"].append(cand_copy)

    for r in root_nodes:
        build_tree_fast(r)

    # Identifica pedidos de venda / requisições
    item_root_demand = {}
    for r in root_nodes:
        if r.get("demand_source_doctype") in ("Sales Order", "Material Request"):
            item_root_demand[r["child_item"]] = (r["demand_source_doctype"], r["demand_source_name"])

    roots_by_demand = defaultdict(list)
    for r in root_nodes:
        doc_type = r.get("demand_source_doctype")
        doc_name = r.get("demand_source_name")

        if (not doc_type or doc_type in ("Manual", "BOM Dependent")) and r["child_item"] in item_root_demand:
            doc_type, doc_name = item_root_demand[r["child_item"]]
            r["demand_source_doctype"] = doc_type
            r["demand_source_name"] = doc_name

        if doc_type and doc_type not in ("Manual", "BOM Dependent"):
            key = (doc_type, doc_name)
        else:
            key = ("Planejamento / Estoque", "Reposição de Segurança e Ponto de Pedido")

        roots_by_demand[key].append(r)

    tree = []
    # 1. Primeiro as ordens com demanda real (Pedidos de Venda, Requisições)
    for (src_type, src_name), branch_roots in roots_by_demand.items():
        if src_type != "Planejamento / Estoque":
            tree.append({
                "is_root_demand": True,
                "demand_source_doctype": src_type,
                "demand_source_name": src_name,
                "label": f"{src_type}: {src_name}",
                "children": branch_roots
            })

    # 2. Depois demandas de manutenção de estoque / segurança
    if ("Planejamento / Estoque", "Reposição de Segurança e Ponto de Pedido") in roots_by_demand:
        branch_roots = roots_by_demand[("Planejamento / Estoque", "Reposição de Segurança e Ponto de Pedido")]
        tree.append({
            "is_root_demand": True,
            "demand_source_doctype": "Estoque de Segurança",
            "demand_source_name": "Reposição Automática (Ponto de Pedido)",
            "label": "Estoque de Segurança: Reposição Automática",
            "children": branch_roots
        })

    return tree

@frappe.whitelist()
def get_ticket_dashboard_metrics(ticket_name):
    """Returns high-level KPI cards data for the workbench."""
    ticket = frappe.get_doc("MRP Ticket", ticket_name)
    
    total_shortages = frappe.db.count("MRP Timeline", {"mrp_ticket": ticket_name, "has_shortage": 1})
    pending_approvals = frappe.db.count("MRP Result", {"mrp_ticket": ticket_name, "status": "Pendente"})
    transfers_count = frappe.db.count("MRP Result", {"mrp_ticket": ticket_name, "supply_type": "Transferência"})
    inconsistencies = frappe.db.count("MRP Log", {"mrp_ticket": ticket_name})
    
    return {
        "status": ticket.status,
        "total_demands": ticket.total_demands or 0,
        "total_suggestions": ticket.total_suggestions or 0,
        "total_produced": ticket.total_produced or 0,
        "total_purchased": ticket.total_purchased or 0,
        "total_transferred": ticket.total_transferred or 0,
        "total_inconsistencies": inconsistencies,
        "pending_approvals": pending_approvals,
        "total_shortages": total_shortages,
        "execution_time": ticket.execution_time_seconds or 0
    }

def style_excel_sheet(ws, title, headers, rows, header_color="1F4E79"):
    """Applies professional enterprise styling to an openpyxl worksheet."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E0'),
        right=Side(style='thin', color='CBD5E0'),
        top=Side(style='thin', color='CBD5E0'),
        bottom=Side(style='thin', color='CBD5E0')
    )

    # Title Banner
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = title
    title_cell.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="102A43", end_color="102A43", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 28

    # Header Row
    ws.append([]) # row 2 spacer
    ws.row_dimensions[2].height = 6

    ws.append(headers)
    ws.row_dimensions[3].height = 24

    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=3, column=col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border

    # Data Rows
    current_row = 4
    for r in rows:
        ws.append(r)
        fill = zebra_fill if (current_row % 2 == 0) else white_fill
        ws.row_dimensions[current_row].height = 20

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.fill = fill
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)

            val = cell.value
            if isinstance(val, (int, float)):
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        current_row += 1

    # Auto Column Widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row > 2 and cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

@frappe.whitelist()
def download_manual_demand_template():
    """Generates and downloads a clean Excel template for importing manual demands."""
    import openpyxl
    import io
    from frappe.desk.utils import provide_binary_file

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Demanda Manual"

    title = "ERPZ MRP — Modelo de Importação de Demanda Manual"
    headers = [
        "Código do Item*", "Quantidade*", "Data da Necessidade (DD/MM/AAAA)*",
        "Depósito", "Observação / Justificativa"
    ]
    rows = [
        ["TEST-MRP-A", 150.0, "25/10/2026", "Mercadorias Em Trânsito - AD", "Demanda extra comercial"],
        ["TEST-MRP-C", 50.0, "28/10/2026", "Mercadorias Em Trânsito - AD", "Reposição extraordinária"]
    ]

    style_excel_sheet(ws, title, headers, rows, header_color="2B6CB0")

    buf = io.BytesIO()
    wb.save(buf)
    provide_binary_file("Modelo_Importacao_Demanda_Manual_MRP", "xlsx", buf.getvalue())

@frappe.whitelist()
def import_manual_demands(ticket_name, file_url=None):
    """
    Imports manual demands from an uploaded Excel spreadsheet into tabMRP Demand.
    """
    import openpyxl
    import io

    if not frappe.has_permission("MRP Ticket", "write"):
        frappe.throw(_("Sem permissão para importar demandas neste Ticket."))

    ticket = frappe.get_doc("MRP Ticket", ticket_name)
    primary_wh = frappe.db.get_value("Warehouse", {"company": ticket.company, "is_group": 0}, "name")

    file_content = None
    if "file" in frappe.request.files:
        file_content = frappe.request.files["file"].read()
    elif file_url:
        _file = frappe.get_doc("File", {"file_url": file_url})
        file_content = _file.get_content()

    if not file_content:
        frappe.throw(_("Nenhum arquivo enviado para importação."))

    wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
    ws = wb.active

    imported = 0
    errors = []

    # Iterate rows starting from row 4 (since rows 1-3 are title, spacer, and headers)
    # If file was made without banner, detect header row
    header_row_idx = 1
    for r_idx in range(1, 10):
        val = str(ws.cell(row=r_idx, column=1).value or "").lower()
        if "código" in val or "codigo" in val or "item" in val:
            header_row_idx = r_idx
            break

    start_data_row = header_row_idx + 1

    for row_num in range(start_data_row, ws.max_row + 1):
        item_code = str(ws.cell(row=row_num, column=1).value or "").strip()
        if not item_code:
            continue

        raw_qty = ws.cell(row=row_num, column=2).value
        raw_date = ws.cell(row=row_num, column=3).value
        warehouse = str(ws.cell(row=row_num, column=4).value or "").strip()
        notes = str(ws.cell(row=row_num, column=5).value or "").strip()

        # Validate item
        if not frappe.db.exists("Item", item_code):
            errors.append(f"Linha {row_num}: Item '{item_code}' não encontrado no cadastro do ERPZ.")
            continue

        # Validate qty
        try:
            qty = flt(raw_qty)
            if qty <= 0:
                raise ValueError()
        except Exception:
            errors.append(f"Linha {row_num}: Quantidade inválida '{raw_qty}'.")
            continue

        # Validate date
        demand_date = None
        try:
            if hasattr(raw_date, "strftime"):
                demand_date = getdate(raw_date)
            elif isinstance(raw_date, str) and "/" in raw_date:
                parts = raw_date.strip().split("/")
                if len(parts) == 3:
                    demand_date = getdate(f"{parts[2]}-{parts[1]}-{parts[0]}")
            else:
                demand_date = getdate(raw_date)
        except Exception:
            errors.append(f"Linha {row_num}: Data inválida '{raw_date}'.")
            continue

        item_doc = frappe.db.get_value("Item", item_code, ["item_name", "stock_uom"], as_dict=True)

        # Insert manual demand
        d = frappe.new_doc("MRP Demand")
        d.mrp_ticket = ticket_name
        d.item_code = item_code
        d.item_name = item_doc.item_name or item_code
        d.company = ticket.company
        d.warehouse = warehouse or primary_wh
        d.demand_date = demand_date
        d.quantity = qty
        d.uom = item_doc.stock_uom
        d.source_type = "Manual"
        d.source_name = f"Demanda Manual ({notes})" if notes else "Demanda Manual"
        d.origin_item = item_code
        d.bom_level = 0
        d.insert(ignore_permissions=True)
        imported += 1

    frappe.db.commit()

    return {
        "success": True,
        "imported_count": imported,
        "errors": errors
    }

@frappe.whitelist()
def export_mrp_excel(ticket_name):
    """
    Exports a comprehensive, formatted multi-tab Excel spreadsheet:
    - Aba 1: Necessidades e Sugestões (formatted table with status, lot, origins, generated docs)
    - Aba 2: Linha do Tempo de Saldos (daily projected evolution)
    - Aba 3: Demandas Consideradas (including Demanda Manual)
    """
    import openpyxl
    import io
    from frappe.desk.utils import provide_binary_file

    wb = openpyxl.Workbook()

    # 1. Sheet: Necessidades e Sugestões
    ws1 = wb.active
    ws1.title = "Necessidades e Sugestões"

    results = frappe.get_all(
        "MRP Result",
        filters={"mrp_ticket": ticket_name},
        fields=[
            "item_code", "item_name", "company", "warehouse", "need_date",
            "supply_date", "gross_demand", "initial_stock", "planned_inflows",
            "projected_balance", "safety_stock", "min_stock", "reorder_point",
            "critical_threshold", "shortage_type", "net_requirement", "suggested_qty",
            "min_order_qty", "supply_type", "origin_item", "origin_name",
            "generated_docname", "status", "situation"
        ],
        order_by="need_date ASC, item_code ASC"
    )

    title1 = f"ERPZ MRP — Resumo de Necessidades e Sugestões de Abastecimento | Ticket: {ticket_name}"
    headers1 = [
        "Código Item", "Descrição do Produto", "Empresa", "Depósito", "Data Necessidade",
        "Data Início", "Demanda Bruta", "Estoque Inicial", "Entradas Previstas",
        "Saldo Projetado", "Estoque Mínimo", "Ponto de Pedido", "Estoque Segurança",
        "Gatilho Crítico (PP+Seg)", "Tipo de Ruptura", "Necessidade Líquida", "Qtd Sugerida",
        "Lote Mínimo", "Tipo Abastecimento", "Produto Originador", "Origem / Demanda",
        "Doc Gerado (OP/OC)", "Situação / Motivo"
    ]
    rows1 = []
    for r in results:
        rows1.append([
            r.item_code, r.item_name, r.company, r.warehouse,
            str(r.need_date) if r.need_date else "",
            str(r.supply_date) if r.supply_date else "",
            flt(r.gross_demand), flt(r.initial_stock), flt(r.planned_inflows),
            flt(r.projected_balance), flt(r.min_stock), flt(r.reorder_point),
            flt(r.safety_stock), flt(r.critical_threshold), r.shortage_type or "Normal",
            flt(r.net_requirement), flt(r.suggested_qty), flt(r.min_order_qty), r.supply_type,
            r.origin_item or "", r.origin_name or "", r.generated_docname or "Pendente",
            r.situation or ""
        ])
    style_excel_sheet(ws1, title1, headers1, rows1, header_color="1F4E79")

    # 2. Sheet: Linha do Tempo de Saldos
    ws2 = wb.create_sheet(title="Linha do Tempo de Saldos")
    timelines = frappe.get_all(
        "MRP Timeline",
        filters={"mrp_ticket": ticket_name},
        fields=[
            "timeline_date", "item_code", "item_name", "company", "initial_balance",
            "inflows", "outflows", "projected_balance", "safety_stock",
            "min_stock", "reorder_point", "critical_threshold", "shortage_type",
            "shortage_qty", "suggested_inflow", "has_shortage"
        ],
        order_by="timeline_date ASC, item_code ASC"
    )
    title2 = f"ERPZ MRP — Linha do Tempo Cronológica do Saldo Projetado | Ticket: {ticket_name}"
    headers2 = [
        "Data", "Código Item", "Descrição", "Empresa", "Saldo Inicial",
        "Entradas Totais", "Saídas (Demandas)", "Saldo Projetado",
        "Estoque Mínimo", "Ponto Pedido + Seg", "Estoque Segurança",
        "Insuficiência / Falta", "Entrada Sugerida", "Situação de Ruptura"
    ]
    rows2 = []
    for t in timelines:
        rows2.append([
            str(t.timeline_date), t.item_code, t.item_name, t.company,
            flt(t.initial_balance), flt(t.inflows + t.suggested_inflow),
            flt(t.outflows), flt(t.projected_balance),
            flt(t.min_stock), flt(t.critical_threshold), flt(t.safety_stock),
            flt(t.shortage_qty), flt(t.suggested_inflow),
            t.shortage_type if t.has_shortage else "OK"
        ])
    style_excel_sheet(ws2, title2, headers2, rows2, header_color="2B6CB0")

    # 3. Sheet: Demandas Consideradas (incluindo Demanda Manual)
    ws3 = wb.create_sheet(title="Demandas Consideradas")
    demands = frappe.get_all(
        "MRP Demand",
        filters={"mrp_ticket": ticket_name},
        fields=["item_code", "item_name", "company", "warehouse", "demand_date", "quantity", "uom", "source_type", "source_name", "origin_item"],
        order_by="demand_date ASC"
    )
    title3 = f"ERPZ MRP — Demandas Consideradas no Cálculo | Ticket: {ticket_name}"
    headers3 = [
        "Código Item", "Descrição", "Empresa", "Depósito", "Data da Demanda",
        "Quantidade", "Unidade", "Tipo de Demanda", "Documento / Origem", "Produto Originador"
    ]
    rows3 = []
    for d in demands:
        rows3.append([
            d.item_code, d.item_name, d.company, d.warehouse, str(d.demand_date),
            flt(d.quantity), d.uom, d.source_type, d.source_name, d.origin_item
        ])
    style_excel_sheet(ws3, title3, headers3, rows3, header_color="2C5282")

    buf = io.BytesIO()
    wb.save(buf)
    provide_binary_file(f"Planejamento_MRP_{ticket_name}", "xlsx", buf.getvalue())

# -------------------------------------------------------------------------
# CENTRAL DE IMPORTAÇÃO DE CADASTROS (PRODUTOS, ESTOQUE, RECURSOS, BOM, CLIENTES, FORNECEDORES)
# -------------------------------------------------------------------------

@frappe.whitelist()
def download_cadastros_template(import_type):
    """Generates and downloads styled Excel template for any of the 7 cadastros."""
    import io
    from frappe.desk.utils import provide_binary_file
    from erpz_mrp.importer import generate_template_workbook

    wb, filename = generate_template_workbook(import_type)
    buf = io.BytesIO()
    wb.save(buf)
    provide_binary_file(filename, "xlsx", buf.getvalue())

@frappe.whitelist()
def import_cadastros(import_type, file_url=None, company=None):
    """
    Imports cadastros from uploaded Excel/CSV file:
    - items: Produtos / Itens (SB1)
    - stock: Estoque dos Itens (SB9)
    - workstations: Recursos / Postos de Trabalho (SH1)
    - alternative_resources: Recursos Alternativos (SH2)
    - bom: Estrutura BOM dos Itens (SG1 - Dono e Componente lado a lado)
    - customers: Clientes (SA1)
    - suppliers: Fornecedores (SA2)
    """
    from erpz_mrp.importer import execute_import

    file_content = None
    filename = "upload.xlsx"

    if "file" in frappe.request.files:
        f = frappe.request.files["file"]
        filename = f.filename
        file_content = f.read()
    elif file_url:
        _file = frappe.get_doc("File", {"file_url": file_url})
        filename = _file.file_name or "upload.xlsx"
        file_content = _file.get_content()

    if not file_content:
        frappe.throw(_("Nenhum arquivo enviado para importação."))

    res = execute_import(import_type, file_content, filename, company)
    return res

@frappe.whitelist()
def get_import_center_options():
    """Returns available cadastros and company list for the frontend modal."""
    companies = frappe.get_all("Company", fields=["name", "company_name", "default_currency"])
    default_company = frappe.defaults.get_user_default("Company") or (companies[0].name if companies else "")

    cadastros = [
        {
            "id": "items",
            "label": "1 - Produtos / Itens",
            "description": "Cadastro e atualização de Itens no ERPZ com estoque mínimo, ponto de pedido, estoque de segurança e lote mínimo.",
            "template_name": "Modelo_Importacao_Produtos_Itens.xlsx"
        },
        {
            "id": "stock",
            "label": "2 - Saldos de Estoque dos Itens",
            "description": "Carga de saldo inicial de estoque por armazém e valor de avaliação via Reconciliação de Estoque no ERPZ.",
            "template_name": "Modelo_Importacao_Estoque_Itens.xlsx"
        },
        {
            "id": "workstations",
            "label": "3 - Recursos / Postos de Trabalho",
            "description": "Cadastro de Postos de Trabalho no ERPZ e Recursos no APS com capacidade diária, eficiência e finais de semana.",
            "template_name": "Modelo_Importacao_Recursos_Postos_Trabalho.xlsx"
        },
        {
            "id": "alternative_resources",
            "label": "4 - Recursos Alternativos",
            "description": "Vinculação de Recursos Alternativos e Secundários aos Postos Principais com prioridade e fator de eficiência.",
            "template_name": "Modelo_Importacao_Recursos_Alternativos.xlsx"
        },
        {
            "id": "bom",
            "label": "5 - Estrutura de Produtos (BOM)",
            "description": "Estrutura do produto onde o Dono do Registro (Pai) e o Componente (Filho) ficam lado a lado por linha, gerando BOMs no ERPZ.",
            "template_name": "Modelo_Importacao_Estrutura_BOM.xlsx"
        },
        {
            "id": "customers",
            "label": "6 - Clientes",
            "description": "Importação de Clientes no ERPZ com Razão Social, Nome Fantasia, CNPJ/CPF, Inscrição Estadual, telefone e e-mail.",
            "template_name": "Modelo_Importacao_Clientes.xlsx"
        },
        {
            "id": "suppliers",
            "label": "7 - Fornecedores",
            "description": "Importação de Fornecedores no ERPZ com Razão Social, CNPJ/CPF, Inscrição Estadual, contatos e endereço.",
            "template_name": "Modelo_Importacao_Fornecedores.xlsx"
        }
    ]

    return {
        "cadastros": cadastros,
        "companies": companies,
        "default_company": default_company
    }

